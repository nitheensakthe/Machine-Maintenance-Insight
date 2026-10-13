import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

os.makedirs('presentation', exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette: Corporate Industrial Navy & Teal
NAVY = RGBColor(31, 78, 120)     # Primary #1F4E78
DARK = RGBColor(38, 38, 38)      # Text #262626
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(41, 128, 185)    # Accent #2980B9
GRAY = RGBColor(242, 244, 247)   # Card bg
AMBER = RGBColor(217, 95, 2)     # Warning #D95F02
LIGHT_BLUE = RGBColor(217, 225, 242)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text="MECHACARE INDUSTRIAL SYSTEMS | DATA SCIENCE ASSESSMENT 2"):
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = TEAL
    
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

# -------------------------------------------------------------
# SLIDE 1: Title Slide
# -------------------------------------------------------------
slide1 = prs.slides.add_slide(blank_layout)

# Background Accent Box
bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()

# Title text
title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.0))
tf = title_box.text_frame
tf.word_wrap = True

p_sub0 = tf.paragraphs[0]
p_sub0.text = "U23AD491 - DATA SCIENCE SPRINT | SKILL BASED ASSESSMENT 2"
p_sub0.font.size = Pt(13)
p_sub0.font.bold = True
p_sub0.font.color.rgb = RGBColor(170, 210, 245)

p_title = tf.add_paragraph()
p_title.text = "Machine Maintenance Insight\nFor MechaCare Industrial Systems"
p_title.font.size = Pt(36)
p_title.font.bold = True
p_title.font.color.rgb = WHITE
p_title.space_before = Pt(15)

p_quest = tf.add_paragraph()
p_quest.text = "Use Case 25: Operating Patterns, Sensor Degradation Signatures & Equipment Prioritization"
p_quest.font.size = Pt(16)
p_quest.font.color.rgb = RGBColor(220, 230, 242)
p_quest.space_before = Pt(15)

p_author = tf.add_paragraph()
p_author.text = "Author: Nitheen Sakthe P E | Role: Lead Data Scientist & Maintenance Reliability Analyst\nDatasets: sensor.csv, machines.csv, usage.csv, alarms.csv, maintenance.csv (Q1 2026 Fleet Logs)"
p_author.font.size = Pt(12)
p_author.font.color.rgb = RGBColor(190, 205, 225)
p_author.space_before = Pt(25)

# -------------------------------------------------------------
# SLIDE 2: Executive Summary & Fleet Metrics
# -------------------------------------------------------------
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Executive Summary: The Industrial Reliability Landscape")

# KPI Cards
kpis = [
    ("30 Machines", "Monitored Fleet Size across 5 Equipment Categories"),
    ("64,016 Hours", "High-frequency Sensor & Usage Synchronized Grain"),
    ("2,510 Alarms", "Cleaned & Standardized (49.2% Critical Severity)"),
    ("268.5 Hours", "Total Unplanned Maintenance Downtime in Q1 2026")
]

for i, (k_val, k_desc) in enumerate(kpis):
    x = Inches(0.8 + i * 2.95)
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(2.8), Inches(1.3))
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT_BLUE
    card.line.color.rgb = TEAL
    
    tf = card.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = k_val
    p0.font.size = Pt(20)
    p0.font.bold = True
    p0.font.color.rgb = NAVY
    p0.alignment = PP_ALIGN.CENTER
    
    p1 = tf.add_paragraph()
    p1.text = k_desc
    p1.font.size = Pt(9.5)
    p1.font.color.rgb = DARK
    p1.alignment = PP_ALIGN.CENTER

# Summary Text Columns
left_box = slide2.shapes.add_textbox(Inches(0.8), Inches(3.1), Inches(5.8), Inches(3.8))
tf_l = left_box.text_frame
tf_l.word_wrap = True
p = tf_l.paragraphs[0]
p.text = "Strategic Business Problem:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_left = [
    ("Core Question: ", "What operating patterns indicate industrial machines that require proactive maintenance attention before catastrophic failure?"),
    ("Fleet Scale: ", "MechaCare monitors 30 industrial assets: CNC Lathes (10), Hydraulic Presses (7), Conveyors (6), Compressors (5), and Milling Machines (2)."),
    ("Data Cleaning Justification: ", "Resolved 35 duplicate sensor rows, 25 duplicate usage rows, imputed missing values using machine-specific medians, and mapped 7 inconsistent alarm strings into 3 engineering categories.")
]
for b_title, b_text in bullets_left:
    p = tf_l.add_paragraph()
    p.text = f"• {b_title}{b_text}"
    p.font.size = Pt(11)
    p.space_before = Pt(8)
    p.font.color.rgb = DARK

right_box = slide2.shapes.add_textbox(Inches(6.9), Inches(3.1), Inches(5.6), Inches(3.8))
tf_r = right_box.text_frame
tf_r.word_wrap = True
p = tf_r.paragraphs[0]
p.text = "Key Data-Driven Insights:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_right = [
    ("Equipment Skew: ", "Hydraulic Presses represent only 23% of the fleet but generate 47.9% of all plant alarms (1,202) and suffer 11.89 hours average downtime."),
    ("The Load Paradox: ", "Operating load has virtually ZERO correlation (r = -0.142) with machine downtime. Machine age, mechanical vibration, and heat buildup are the actual failure drivers."),
    ("Pre-Service Trajectory: ", "Sensor vibration exhibits a statistically significant upward climb (+11.8%) starting 7 days prior to service events, peaking at 2.85 mm/s within 24 hours of intervention.")
]
for b_title, b_text in bullets_right:
    p = tf_r.add_paragraph()
    p.text = f"• {b_title}{b_text}"
    p.font.size = Pt(11)
    p.space_before = Pt(8)
    p.font.color.rgb = DARK

# -------------------------------------------------------------
# SLIDE 3: Equipment Vulnerability & The Load Paradox
# -------------------------------------------------------------
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Analytical Questions 1, 4 & 5: Sensor Extremes, Load vs Downtime")

# Insert Fig 1 & Fig 4
if os.path.exists('figures/01_sensor_distributions.png'):
    slide3.shapes.add_picture('figures/01_sensor_distributions.png', Inches(0.8), Inches(1.5), width=Inches(6.6))
if os.path.exists('figures/04_operating_load_vs_downtime.png'):
    slide3.shapes.add_picture('figures/04_operating_load_vs_downtime.png', Inches(7.5), Inches(1.5), width=Inches(5.0))

notes_box = slide3.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.8))
tf_n = notes_box.text_frame
tf_n.word_wrap = True

p = tf_n.paragraphs[0]
p.text = "Key Evidentiary Findings from Actual Dataset:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_s3 = [
    ("Q1 (Abnormal Ranges): ", "MCH-004 exhibited the highest vibration peak plant-wide (17.20 mm/s, over 3.9x normal threshold) and 1,580 warning hours. MCH-021 registered extreme thermal stress peaking at 95.96°C."),
    ("Q4 (Load Paradox): ", "Average operating load across all machines remains constant (~65.0%). Machines in the low downtime tier average 65.03% load, while high downtime machines average 65.02%. Load does not explain failure."),
    ("Q5 (Machine Types): ", "Hydraulic Presses average 11.89h downtime/machine. CNC Lathes require the highest service count (20 services = 2.0/machine). Conveyors have lowest downtime per event (3.46h) but frequent friction alarms.")
]
for b_title, b_text in bullets_s3:
    p = tf_n.add_paragraph()
    p.text = f"• {b_title}{b_text}"
    p.font.size = Pt(10.5)
    p.space_before = Pt(4)

# -------------------------------------------------------------
# SLIDE 4: Sensor Degradation Signature & Alarm Analysis
# -------------------------------------------------------------
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Analytical Questions 2 & 3: Failure Signatures & Alarm Frequency")

if os.path.exists('figures/02_sensor_degradation_before_maintenance.png'):
    slide4.shapes.add_picture('figures/02_sensor_degradation_before_maintenance.png', Inches(0.8), Inches(1.5), width=Inches(5.6))
if os.path.exists('figures/03_alarm_frequency_and_severity.png'):
    slide4.shapes.add_picture('figures/03_alarm_frequency_and_severity.png', Inches(6.6), Inches(1.5), width=Inches(5.9))

notes_box = slide4.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.9))
tf_n = notes_box.text_frame
tf_n.word_wrap = True

p = tf_n.paragraphs[0]
p.text = "Operational Degradation Signatures & Cleansing Evidence:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_s4 = [
    ("Q2 (Pre-Service Trajectory): ", "Plant vibration baseline averages 2.55 mm/s. 7 days prior to service, it climbs to 2.79 mm/s; at 48h-24h it hits 2.84 mm/s; and within 24h of failure it peaks at 2.85 mm/s (+11.8% elevation). Temperature similarly rises from 58.2°C to 59.1°C."),
    ("Q3 (Alarm Normalization): ", "Standardized 7 inconsistent sensor strings into 3 root types: High Temperature (1,084 events, 43.2%), High Vibration (1,080 events, 43.0%), and Current Spikes (346 events, 13.8%)."),
    ("Severity Distribution: ", "Critical severity accounts for 49.2% (1,236 alarms) vs 50.8% Warning (1,274 alarms). High temperature critical events represent the primary emergency shutdown trigger.")
]
for b_title, b_text in bullets_s4:
    p = tf_n.add_paragraph()
    p.text = f"• {b_title}{b_text}"
    p.font.size = Pt(10.5)
    p.space_before = Pt(4)

# -------------------------------------------------------------
# SLIDE 5: Repeat Failure Analysis & Chronic Assets
# -------------------------------------------------------------
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Analytical Question 6: Chronic Machines & Repeat Failure Patterns")

if os.path.exists('figures/06_repeat_failure_patterns.png'):
    slide5.shapes.add_picture('figures/06_repeat_failure_patterns.png', Inches(0.8), Inches(1.5), width=Inches(7.2))

# Machine Box on right
card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.3), Inches(1.5), Inches(4.2), Inches(3.6))
card.fill.solid()
card.fill.fore_color.rgb = GRAY
card.line.color.rgb = TEAL

tf_c = card.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "Top Repeat Failure Assets (3 Services):"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = NAVY

repeat_details = [
    ("MCH-005 (CNC Lathe, 5y): ", "10.31h downtime. Interventions: Calibration -> Lubrication -> Inspection. Lack of component overhaul allows alignment drift to recur."),
    ("MCH-016 (Conveyor, 9y): ", "13.82h downtime. Progressive failure: Lubrication starved -> Motor burned -> Bearing replaced. Clear cascading wear."),
    ("MCH-022 (CNC Lathe, 1y): ", "12.54h downtime. Infant mortality: 2 recalibrations and 1 lubrication in first 90 days. Indicates commissioning / assembly defects."),
    ("MCH-028 (CNC Lathe, 9y): ", "12.62h downtime. Cascading wear: Headstock bearing replacement followed by lubrication and motor repair.")
]
for r_title, r_desc in repeat_details:
    p = tf_c.add_paragraph()
    p.text = f"• {r_title}{r_desc}"
    p.font.size = Pt(9.5)
    p.space_before = Pt(4)

notes_box = slide5.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.8))
tf_n = notes_box.text_frame
tf_n.word_wrap = True

p = tf_n.paragraphs[0]
p.text = "Strategic Engineering Lessons from Repeat Failures:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = NAVY

bullets_s5 = [
    ("Infant Mortality vs Wear-out: ", "Chronic failures occur at opposite ends of asset lifecycle: Brand-new machines (MCH-022, age 1y) suffer assembly/alignment bedding issues, while 9-year assets (MCH-016, MCH-028) experience mechanical cascading fatigue."),
    ("Preventing Cascading Damage: ", "MCH-016 proves that insufficient initial lubrication creates downstream motor overheating and bearing seizure, tripling total repair cost.")
]
for b_title, b_text in bullets_s5:
    p = tf_n.add_paragraph()
    p.text = f"• {b_title}{b_text}"
    p.font.size = Pt(10.5)
    p.space_before = Pt(4)

# -------------------------------------------------------------
# SLIDE 6: Feature Engineering, PCA & Predictive Modeling Limitations
# -------------------------------------------------------------
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "Feature Engineering, Modeling Rigor & Predictive Limitations")

col1 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.5))
tf1 = col1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "Feature Engineering & Selection:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = NAVY

fe_points = [
    ("Rolling Sensor Dynamics: ", "Engineered 7-day rolling mean, max vibration, and rolling temperature to capture mechanical degradation slope without lookahead bias."),
    ("Event Memory: ", "Created cumulative sensor warning hours (sensor_warning), historical service counts, and prior downtime hours per machine."),
    ("Correlation & PCA Assessment: ", "Multicollinearity observed between vibration_mean and rolling vibration (r > 0.85). PCA was evaluated but NOT recommended for production deployment because preserving physical interpretability (mm/s, °C) is vital for field technicians.")
]
for title, text in fe_points:
    p = tf1.add_paragraph()
    p.text = f"• {title}{text}"
    p.font.size = Pt(10.5)
    p.space_before = Pt(6)

col2 = slide6.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.8), Inches(5.5))
tf2 = col2.text_frame
tf2.word_wrap = True

p = tf2.paragraphs[0]
p.text = "Model Validation & The Class Imbalance Reality:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = NAVY

model_points = [
    ("Chronological Split: ", "Trained on Jan 1 - Feb 28, 2026 (1,770 machine-days); evaluated out-of-time on March 1 - March 31, 2026 (930 machine-days) to strictly prevent data leakage."),
    ("Model Comparison: ", "Evaluated Stratified Baseline vs Balanced Random Forest (100 trees, max_depth=5). Baseline F1 = 0.11; Random Forest F1 = 0.18, Precision = 13.1%, Recall = 28.3%, ROC-AUC = 0.534, PR-AUC = 0.115."),
    ("Why High Accuracy is Deceptive: ", "Overall accuracy appears high (74%), but this is purely driven by class imbalance (negative class = 90.1% of test set). High accuracy is a statistical illusion when failures are rare."),
    ("Critical Scientific Distinction: ", "With only 51 historical service events across 90 days, machine learning cannot produce a clinically reliable failure forecast. MechaCare must rely on our Rule-Based Multi-Sensor Health Score for immediate scheduling.")
]
for title, text in model_points:
    p = tf2.add_paragraph()
    p.text = f"• {title}{text}"
    p.font.size = Pt(10.5)
    p.space_before = Pt(6)

# -------------------------------------------------------------
# SLIDE 7: Business Recommendations & Prioritization Framework
# -------------------------------------------------------------
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Prioritization Framework & Strategic Business Recommendations")

# 3 Priority Columns
priorities = [
    ("Tier 1: Immediate Overhaul (Days 1-7)", [
        ("MCH-004 (Hydraulic Press): ", "Plant-worst 21.5h downtime & 457 alarms. Overhaul main hydraulic pump, relief valve & seals."),
        ("MCH-014 (Compressor): ", "19.2h downtime, severe thermal spikes. Clean intercooler & replace discharge reed valves."),
        ("MCH-015 (CNC Lathe): ", "16.7h downtime, 10.3 mm/s vibration. Conduct laser spindle alignment & bearing runout check.")
    ], AMBER),
    ("Tier 2: Lubrication & Repeat-Failure Program", [
        ("Conveyor Fleet Reliability: ", "Implement automatic greasing system on MCH-016 & MCH-025 to stop lubricant starvation from destroying motors."),
        ("Commissioning Quality Gate: ", "Audit newly commissioned assets (MCH-022, 1y old) with mandatory 100-hour vibration & torque verification."),
        ("Downtime Reduction Goal: ", "Eliminates ~65 hours of chronic recurring downtime in the next quarter.")
    ], TEAL),
    ("Tier 3: Operational ROI & Architecture", [
        ("Telemetry Thresholds: ", "Deploy real-time automated alerts at 4.40 mm/s vibration and 78.0°C temperature with automated work orders."),
        ("Sensor Instrumentation: ", "Add IoT acoustic sensors to Hydraulic Presses to detect cavitation before pump failure occurs."),
        ("Projected Business ROI: ", "Preventing 120 hours of unplanned downtime saves MechaCare an estimated $180,000+ annually.")
    ], NAVY)
]

for i, (p_title, p_items, p_color) in enumerate(priorities):
    x = Inches(0.8 + i * 3.95)
    card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(3.8), Inches(5.4))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = p_color
    card.line.width = Pt(2)
    
    tf = card.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = p_title
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = p_color
    p0.alignment = PP_ALIGN.CENTER
    
    for item_head, item_body in p_items:
        p = tf.add_paragraph()
        p.text = f"• {item_head}{item_body}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

prs.save('presentation/MechaCare_Machine_Maintenance_Presentation.pptx')
print('Successfully saved presentation/MechaCare_Machine_Maintenance_Presentation.pptx!')
